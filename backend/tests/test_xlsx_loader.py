import io
import pytest
import openpyxl
from backend.app.services.ingestion.loaders.xlsx_loader import XLSXLoader

def create_sample_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    # Sheet 1: Students
    ws1 = wb.active
    ws1.title = "Students"
    ws1.append(["Roll No", "Full Name", "Major", "GPA"])
    ws1.append(["CS101", "Alex Chen", "Computer Science", 3.9])
    ws1.append(["EE202", "Maya Patel", "Electrical Engineering", 3.7])

    # Sheet 2: Courses
    ws2 = wb.create_sheet(title="Courses")
    ws2.append(["Course Code", "Course Name", "Credits"])
    ws2.append(["CS501", "Advanced Database Systems", 4])
    ws2.append(["CS502", "Distributed Computing", 4])

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()

def test_xlsx_loader_multi_sheet():
    xlsx_bytes = create_sample_xlsx_bytes()
    loader = XLSXLoader(xlsx_bytes, file_name="academic_records.xlsx")
    doc = loader.load()

    assert doc.title == "academic_records"
    assert doc.source_type == "xlsx"
    assert doc.file_name == "academic_records.xlsx"
    assert len(doc.segments) == 4 # 2 rows in Students + 2 rows in Courses

    # Check sheet names preserved
    sheet_names = {seg.sheet_name for seg in doc.segments}
    assert sheet_names == {"Students", "Courses"}

    # Check row numbers
    student_rows = [seg.row_number for seg in doc.segments if seg.sheet_name == "Students"]
    assert student_rows == [2, 3]

    # Check cell key-value contents
    first_seg = doc.segments[0]
    assert "Full Name: Alex Chen" in first_seg.text
    assert "Major: Computer Science" in first_seg.text

def test_xlsx_loader_empty_error():
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    loader = XLSXLoader(buf.getvalue(), file_name="empty.xlsx")
    with pytest.raises(ValueError, match="no readable data rows"):
        loader.load()
