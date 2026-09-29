import io
import pytest
import docx
from backend.app.services.ingestion.loaders.docx_loader import DOCXLoader

def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.core_properties.title = "Database Systems Guide"
    doc.core_properties.author = "Dr. Stone"

    doc.add_heading("Relational Model", level=1)
    doc.add_paragraph("The relational model represents data as relations or tables composed of rows and columns.")

    doc.add_heading("Normalization", level=2)
    doc.add_paragraph("Normalization minimizes redundancy and dependency by organizing fields and table relations.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Form"
    table.cell(0, 1).text = "Description"
    table.cell(1, 0).text = "1NF"
    table.cell(1, 1).text = "Eliminates duplicate columns from the same table"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_docx_loader_success():
    docx_bytes = create_sample_docx_bytes()
    loader = DOCXLoader(docx_bytes, file_name="dbms_guide.docx")
    doc = loader.load()

    assert doc.title == "Database Systems Guide"
    assert doc.source_type == "docx"
    assert doc.file_name == "dbms_guide.docx"
    assert len(doc.segments) >= 3

    # Check section title tracking
    sections = [seg.section_title for seg in doc.segments]
    assert any("Relational Model" in s for s in sections if s)
    assert any("Normalization" in s for s in sections if s)

    # Check table extraction
    assert any("1NF" in seg.text for seg in doc.segments)

def test_docx_loader_empty_error():
    doc = docx.Document()
    buf = io.BytesIO()
    doc.save(buf)

    loader = DOCXLoader(buf.getvalue(), file_name="empty.docx")
    with pytest.raises(ValueError, match="no readable text"):
        loader.load()
