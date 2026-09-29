import io
import os
from typing import Dict, Any, List, Optional
import openpyxl

from backend.app.services.ingestion.loaders.base import BaseLoader, ExtractedDocument, ExtractedSegment
from backend.app.core.logging import logger

class XLSXLoader(BaseLoader):
    """
    Ingests Excel spreadsheets (.xlsx) across all sheets using openpyxl.
    Preserves sheet names, row numbers, and column headers.
    """

    def __init__(self, file_path_or_bytes: str | bytes, file_name: Optional[str] = None):
        self.file_path_or_bytes = file_path_or_bytes
        self.file_name = file_name or (
            os.path.basename(file_path_or_bytes) if isinstance(file_path_or_bytes, str) else "workbook.xlsx"
        )

    def load(self) -> ExtractedDocument:
        logger.info(f"Parsing Excel spreadsheet: {self.file_name}")

        try:
            if isinstance(self.file_path_or_bytes, bytes):
                wb = openpyxl.load_workbook(io.BytesIO(self.file_path_or_bytes), data_only=True)
            else:
                wb = openpyxl.load_workbook(self.file_path_or_bytes, data_only=True)
        except Exception as e:
            raise ValueError(f"Failed to open Excel file '{self.file_name}': {str(e)}")

        segments: List[ExtractedSegment] = []
        full_text_parts: List[str] = []
        sheet_summaries: Dict[str, int] = {}
        title = os.path.splitext(self.file_name)[0]

        try:
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                # First non-empty row is treated as headers
                header_row_idx = 0
                headers = []
                while header_row_idx < len(rows):
                    candidate_headers = rows[header_row_idx]
                    if any(c is not None and str(c).strip() for c in candidate_headers):
                        headers = [str(c).strip() if c is not None else f"Column_{i+1}" for i, c in enumerate(candidate_headers)]
                        break
                    header_row_idx += 1

                if not headers:
                    continue

                sheet_row_count = 0
                # Process subsequent rows
                for row_idx, row_values in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
                    if not any(v is not None and str(v).strip() for v in row_values):
                        continue

                    row_parts = []
                    for col_idx, cell_value in enumerate(row_values):
                        if cell_value is None:
                            continue
                        str_val = str(cell_value).strip()
                        if not str_val:
                            continue
                        header_name = headers[col_idx] if col_idx < len(headers) else f"Column_{col_idx+1}"
                        row_parts.append(f"{header_name}: {str_val}")

                    if row_parts:
                        row_text = f"[{sheetname}] Row {row_idx}: " + " | ".join(row_parts)
                        full_text_parts.append(row_text)
                        segments.append(
                            ExtractedSegment(
                                text=row_text,
                                sheet_name=sheetname,
                                row_number=row_idx,
                                section_title=f"{sheetname} (Row {row_idx})"
                            )
                        )
                        sheet_row_count += 1

                sheet_summaries[sheetname] = sheet_row_count

            if not segments:
                raise ValueError(f"Excel file '{self.file_name}' contains no readable data rows.")

            return ExtractedDocument(
                title=title,
                source_type="xlsx",
                file_name=self.file_name,
                full_text="\n".join(full_text_parts),
                metadata={
                    "file_name": self.file_name,
                    "sheets": list(sheet_summaries.keys()),
                    "rows_per_sheet": sheet_summaries,
                    "total_rows": len(segments)
                },
                segments=segments
            )
        finally:
            wb.close()
